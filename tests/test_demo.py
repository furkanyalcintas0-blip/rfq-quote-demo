"""Offline checks using temporary synthetic assets and fixed evaluation dates."""

import csv
import hashlib
import io
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from reportlab.pdfgen import canvas

from rfq_quote.cli import main
from rfq_quote.export import export_quote, summarize
from rfq_quote.extract import Item, SourceLine, extract_pdf
from rfq_quote.match import Part, Price, load_catalogue, load_prices, match_items, normalize_code
from samples.make_samples import make_samples


ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 3)
STARTED = datetime(2026, 10, 3, 12, 0, 0, tzinfo=timezone.utc)


class DemoTests(unittest.TestCase):
    """Exercise every public pipeline function with no network or shared output."""

    def setUp(self):
        """Generate isolated deterministic sample files for each test."""
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.files = make_samples(self.root / "samples")
        self.catalogue = load_catalogue(self.files["catalogue.csv"])
        self.prices = load_prices(self.files["price_list.csv"])

    def _item(self, code="TRM-101", description="Terminal block type 1", unit="ea", qty="3"):
        return Item("1", code, description, Decimal(qty), unit, SourceLine(1, 6, "source item"))

    def _matches(self, name="rfq_a.pdf"):
        extraction = extract_pdf(self.files[name])
        return extraction, match_items(extraction.items, self.catalogue, self.prices, TODAY)

    def _export(self, started=STARTED):
        extraction, matches = self._matches()
        return export_quote(matches, extraction.unparsed, self.files["rfq_a.pdf"],
                            self.root / "out", TODAY, started,
                            (self.files["catalogue.csv"], self.files["price_list.csv"]))

    def _pdf(self, name, lines):
        path = self.root / name
        pdf = canvas.Canvas(str(path), invariant=1)
        for i, line in enumerate(lines):
            pdf.drawString(20, 800 - 20 * i, line)
        pdf.save()
        return path

    def _csv(self, name, headers, rows):
        path = self.root / name
        with path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)
        return path

    def test_samples_are_deterministic_and_have_50_parts(self):
        """Regeneration produces the same bytes and the requested catalogue size."""
        before = {name: path.read_bytes() for name, path in self.files.items()}
        regenerated = make_samples(self.root / "samples")
        self.assertEqual(before, {name: path.read_bytes() for name, path in regenerated.items()})
        self.assertEqual(len(self.catalogue), 50)
        self.assertEqual(len(self.prices), 50)

    def test_extract_layout_a(self):
        """The first layout preserves quantities, codes, and source locations."""
        extraction = extract_pdf(self.files["rfq_a.pdf"])
        self.assertEqual(len(extraction.items), 11)
        self.assertEqual(extraction.items[0].part_no, "TRM-101")
        self.assertEqual(extraction.items[1].qty, Decimal("35.5"))
        self.assertEqual(extraction.items[6].part_no, "")
        self.assertEqual(extraction.items[0].source.page, 1)
        self.assertIn("TRM-101", extraction.items[0].source.text)

    def test_extract_layout_b(self):
        """The reordered second layout reads code, quantity and unit correctly."""
        extraction = extract_pdf(self.files["rfq_b.pdf"])
        self.assertEqual(len(extraction.items), 11)
        self.assertEqual(extraction.items[0].part_no, "SWI-104")
        self.assertEqual(extraction.items[0].qty, Decimal("8"))
        self.assertEqual(extraction.items[0].unit, "ea")
        self.assertEqual(len(extraction.unparsed), 1)

    def test_unparsed_item_is_retained(self):
        """An invalid quantity leaves an original row and a reason."""
        extraction = extract_pdf(self.files["rfq_a.pdf"])
        self.assertEqual(len(extraction.unparsed), 1)
        self.assertIn("several", extraction.unparsed[0].source.text)
        self.assertIn("Invalid item fields", extraction.unparsed[0].reason)

    def test_unknown_lines_and_invalid_quantities_are_retained(self):
        """Unknown prose, missing fields, and nonfinite quantities are never lost."""
        header = "Line | Part no | Description | Qty | Unit"
        lines = ["Unexpected preamble", header, "Unexpected footnote"]
        lines += [f"1 | TRM-101 | Terminal block type 1 | {qty} | ea"
                  for qty in ("0", "-1", "NaN", "Infinity", "1,000")]
        lines += ["2 | TRM-101 | | 1 | ea", "BROKEN TABLE ROW"]
        extraction = extract_pdf(self._pdf("bad_rows.pdf", lines))
        self.assertEqual(len(extraction.items), 0)
        self.assertEqual(len(extraction.unparsed), len(lines) - 1)
        self.assertEqual([r.source.text for r in extraction.unparsed], lines[:1] + lines[2:])

    def test_quantity_format_is_explicit_and_ambiguous_rows_survive(self):
        """Only positive integers and decimals without three fractional digits parse."""
        accepted = ("5", "1500", "0005", "1.5", "1.50", "1.5000", "0.01", "0.0010")
        rejected = ("1.500", "12.345", "0.001", "1,5", "1_000", "1e3", "1 000",
                    "+5", "-2", "0", "", "0.00", "NaN", "Infinity", "abc")
        for layout in (("Line", "Part no", "Description", "Qty", "Unit"),
                       ("Line", "Qty", "Unit", "Description", "Part no")):
            with self.subTest(layout=layout):
                rows = []
                for number, qty in enumerate(accepted + rejected, 1):
                    row = dict(Line=str(number), **{"Part no": "TRM-101",
                               "Description": "Terminal block", "Qty": qty, "Unit": "ea"})
                    rows.append(" | ".join(row[column] for column in layout))
                extraction = extract_pdf(self._pdf("quantities.pdf", [" | ".join(layout), *rows]))
                self.assertEqual([item.qty for item in extraction.items],
                                 [Decimal(qty) for qty in accepted])
                # PDF text extraction collapses the padding around empty cells.
                self.assertEqual([tuple(cell.strip() for cell in row.source.text.split("|"))
                                  for row in extraction.unparsed],
                                 [tuple(cell.strip() for cell in row.split("|"))
                                  for row in rows[len(accepted):]])
                self.assertEqual([row.source.line for row in extraction.unparsed],
                                 list(range(len(accepted) + 2, len(rows) + 2)))
                self.assertTrue(all(row.reason for row in extraction.unparsed))

    def test_quantity_reasons_are_readable_in_excel(self):
        """Invalid quantities retain explanatory text in both workbook sheets."""
        quantities = ("1.500", "abc", "", "0")
        expected = ("quantity '1.500' is ambiguous (thousands or decimal?)",
                    "quantity 'abc' is not a number", "quantity is empty",
                    "quantity '0' must be positive")
        path = self._pdf("quantity_reasons.pdf", ["Line | Part no | Description | Qty | Unit",
                         *[f"{i} | TRM-101 | Terminal block | {qty} | ea"
                           for i, qty in enumerate(quantities, 1)]])
        extraction = extract_pdf(path)
        for number, reason in enumerate(expected, 2):
            with self.subTest(reason=reason):
                actual = next((row.reason for row in extraction.unparsed
                               if row.source.line == number), "")
                self.assertIn(reason, actual)
        self.assertEqual(len(extraction.unparsed), len(quantities))
        output = export_quote([], extraction.unparsed, path, self.root / "out", TODAY, STARTED)
        book = load_workbook(output)
        self.addCleanup(book.close)
        for sheet, column in ((book["Quote draft"], 11), (book["Needs review"], 5)):
            reasons = [sheet.cell(i, column).value for i in range(2, 6)]
            self.assertEqual(reasons, [row.reason for row in extraction.unparsed])
            self.assertFalse(any("ConversionSyntax" in reason or "<class" in reason
                                 for reason in reasons))

    def test_readme_documents_standalone_commands_and_sample_limits(self):
        """A standalone reader gets setup commands and honest sample limitations."""
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        for text in ("python -m venv .venv", "pip install -r requirements.txt",
                     "python samples/make_samples.py", "python -m rfq_quote.cli",
                     "python -m unittest discover -s tests", "same generator",
                     "same anomaly categories", "two different column layouts",
                     "exactly three fractional digits", "1.500"):
            with self.subTest(text=text):
                self.assertTrue(text in readme, f"README is missing: {text}")
        self.assertNotIn("../../.venv/bin/python", readme)
        self.assertNotIn("parent workspace", readme)

    def test_temporary_assets_use_system_temp_and_legacy_paths_are_ignored(self):
        """Test assets stay in system temp; old test folders are ignored as well."""
        with self.subTest(check="system temp"):
            self.assertEqual(self.root.parent.resolve(), Path(tempfile.gettempdir()).resolve())
        with self.subTest(check="legacy ignore pattern"):
            self.assertIn("tests/.test-*", (ROOT / ".gitignore").read_text().splitlines())

    def test_image_only_or_unsupported_pdf_is_rejected(self):
        """An empty PDF or a PDF with another layout cannot become a valid draft."""
        for lines in ([], ["Unrecognised table header"]):
            with self.subTest(lines=lines):
                with self.assertRaisesRegex(ValueError, "No supported text table"):
                    extract_pdf(self._pdf("unsupported.pdf", lines))

    def test_multiple_pages_and_repeated_headers(self):
        """Repeated headers do not become unparsed rows and page numbers survive."""
        path = self.root / "multipage.pdf"
        pdf = canvas.Canvas(str(path), invariant=1)
        for page in (1, 2):
            pdf.drawString(20, 800, "Line | Part no | Description | Qty | Unit")
            pdf.drawString(20, 780, f"{page} | TRM-101 | Terminal block type 1 | 1 | ea")
            pdf.showPage()
        pdf.save()
        extraction = extract_pdf(path)
        self.assertEqual([i.source.page for i in extraction.items], [1, 2])
        self.assertEqual(extraction.unparsed, [])

    def test_load_catalogue(self):
        """Catalogue fields retain synthetic descriptions and sales units."""
        self.assertEqual(self.catalogue["CBL-102"], Part("CBL-102", "Control cable type 2", "m"))

    def test_load_prices(self):
        """Price values are decimals and validity dates are parsed."""
        self.assertEqual(self.prices["TRM-101"].unit_price, Decimal("1.98"))
        self.assertEqual(self.prices["REL-101"].valid_until, date(2020, 1, 1))

    def test_invalid_catalogue_is_rejected(self):
        """Bad CSV schemas, missing values and duplicate exact codes fail early."""
        cases = [(["part_no"], [["X"]]),
                 (["part_no", "description", "unit"], [["X", "", "ea"]]),
                 (["part_no", "description", "unit"], [["X", "Part", "ea"]] * 2),
                 (["part_no", "description", "unit"], []),
                 (["part_no", "description", "unit"], [["X", "Part", "ea", "extra"]])]
        for headers, rows in cases:
            with self.subTest(rows=rows):
                with self.assertRaises(ValueError):
                    load_catalogue(self._csv("bad.csv", headers, rows))

    def test_invalid_price_values_are_rejected(self):
        """Invalid numeric values, validity dates and currencies cannot be priced."""
        for value, currency, until in (("NaN", "USD", "2099-12-31"),
                                       ("Infinity", "USD", "2099-12-31"),
                                       ("oops", "USD", "2099-12-31"),
                                       ("1", "usd", "2099-12-31"),
                                       ("1", "USD", "invalid")):
            with self.subTest(value=value, currency=currency, until=until):
                path = self._csv("bad_prices.csv", ["part_no", "unit_price", "currency", "valid_until"],
                                 [["X", value, currency, until]])
                with self.assertRaises(ValueError):
                    load_prices(path)

    def test_duplicate_prices_are_rejected(self):
        """Two approved price entries for the same exact part code are ambiguous."""
        path = self._csv("duplicate.csv", ["part_no", "unit_price", "currency", "valid_until"],
                         [["X", "1", "USD", "2099-12-31"]] * 2)
        with self.assertRaisesRegex(ValueError, "duplicate"):
            load_prices(path)

    def test_normalize_code(self):
        """Normalisation ignores only whitespace, hyphens, and case."""
        self.assertEqual(normalize_code(" trm - 101\t"), "TRM101")
        self.assertNotEqual(normalize_code("TRM_101"), "TRM101")

    def test_exact_match_uses_approved_price(self):
        """Exact codes produce priced rows from the supplied price list."""
        result = match_items([self._item()], self.catalogue, self.prices, TODAY)[0]
        self.assertEqual(result.status, "priced")
        self.assertEqual(result.price, self.prices["TRM-101"])

    def test_normalized_match_is_priced_with_note(self):
        """A unique normalised code gets a price and an explicit audit note."""
        result = match_items([self._item("trm 101")], self.catalogue, self.prices, TODAY)[0]
        self.assertEqual(result.status, "priced")
        self.assertEqual(result.part.part_no, "TRM-101")
        self.assertIn("Normalized code match", result.reason)

    def test_description_only_requires_review_and_top_three(self):
        """Description suggestions never select or price an alternative."""
        result = match_items([self._item("", "Terminal blok type 1")],
                             self.catalogue, self.prices, TODAY)[0]
        self.assertEqual(result.status, "needs review")
        self.assertIsNone(result.price)
        self.assertIsNone(result.part)
        self.assertEqual(len(result.candidates), 3)
        self.assertEqual(result.candidates[0].part.part_no, "TRM-101")
        self.assertEqual([c.score for c in result.candidates],
                         sorted([c.score for c in result.candidates], reverse=True))

    def test_unknown_code_is_not_substituted(self):
        """An unknown explicit code stays unknown even with an exact description."""
        result = match_items([self._item("ZZZ-900")], self.catalogue, self.prices, TODAY)[0]
        self.assertEqual(result.status, "not in catalogue")
        self.assertIsNone(result.price)
        self.assertIsNone(result.part)

    def test_unrelated_description_is_not_in_catalogue(self):
        """A missing code and unrelated description produce no suggested price."""
        result = match_items([self._item("", "zzzzzzzzzz")], self.catalogue, self.prices, TODAY)[0]
        self.assertEqual(result.status, "not in catalogue")
        self.assertEqual(result.candidates, ())

    def test_flux_harmonizer_has_no_unrelated_candidates(self):
        """Shared character fragments do not suggest a fuse for a flux module."""
        result = match_items([self._item("", "Flux harmonizer module")],
                             self.catalogue, self.prices, TODAY)[0]
        self.assertEqual(result.status, "not in catalogue")
        self.assertEqual(result.candidates, ())
        self.assertIsNone(result.price)

    def test_close_description_typo_remains_an_unpriced_candidate(self):
        """A real spelling variant remains reviewable without selecting or pricing."""
        part = Part("CB-16", "Circuit breaker 16A", "ea")
        catalogue = {**self.catalogue, part.part_no: part}
        prices = {**self.prices, part.part_no: Price(part.part_no, Decimal("10"), "USD", TODAY)}
        result = match_items([self._item("", "Circut breaker 16A")], catalogue, prices, TODAY)[0]
        self.assertEqual(result.status, "needs review")
        self.assertEqual(result.candidates[0].part, part)
        self.assertIsNone(result.part)
        self.assertIsNone(result.price)
        output = export_quote([result], [], self.files["rfq_a.pdf"], self.root / "out", TODAY, STARTED)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertIsNone(book.active.cell(2, 7).value)
        self.assertIsNone(book.active.cell(2, 9).value)

    def test_nonpositive_csv_prices_are_reviewed_without_stopping_valid_rows(self):
        """Zero and negative CSV prices survive for review while good rows price."""
        for value in ("0", "-1", "-0.01"):
            with self.subTest(value=value):
                path = self._csv("review_prices.csv",
                                 ["part_no", "unit_price", "currency", "valid_until"],
                                 [["TRM-101", value, "USD", "2099-12-31"],
                                  ["TRM-102", "2", "USD", "2099-12-31"]])
                prices = load_prices(path)
                matches = match_items([self._item(), self._item("TRM-102")],
                                      self.catalogue, prices, TODAY)
                self.assertEqual([m.status for m in matches], ["invalid price", "priced"])
                self.assertIn("must be positive", matches[0].reason)
                self.assertIsNone(matches[0].price)
                self.assertEqual(summarize(matches, [])["needs review"], 1)
                output = export_quote(matches, [], self.files["rfq_a.pdf"],
                                      self.root / "out", TODAY, STARTED)
                book = load_workbook(output)
                self.addCleanup(book.close)
                for column in (7, 8, 9):
                    self.assertIsNone(book.active.cell(2, column).value)
                self.assertEqual(book.active.cell(2, 10).value, "invalid price")
                self.assertEqual(book["Needs review"].cell(2, 4).value, "invalid price")
                self.assertEqual(dict(book["Run info"].iter_rows(min_row=2, values_only=True))
                                 ["needs review"], 1)

    def test_invalid_direct_prices_cannot_bypass_matching_guard(self):
        """Direct price objects are checked, even when their validity has expired."""
        for value in ("0", "-2", "NaN", "Infinity", "-Infinity"):
            with self.subTest(value=value):
                prices = {"TRM-101": Price("TRM-101", Decimal(value), "USD", date(2020, 1, 1))}
                result = match_items([self._item()], self.catalogue, prices, TODAY)[0]
                self.assertEqual(result.status, "invalid price")
                self.assertIsNone(result.price)

    def test_price_expired_and_validity_boundary(self):
        """Expiry uses an injected date and includes the final validity day."""
        prices = {"TRM-101": Price("TRM-101", Decimal("2.50"), "USD", TODAY)}
        self.assertEqual(match_items([self._item()], self.catalogue, prices, TODAY)[0].status, "priced")
        expired = match_items([self._item()], self.catalogue, prices, date(2026, 10, 4))[0]
        self.assertEqual(expired.status, "price expired")
        self.assertIsNone(expired.price)

    def test_missing_approved_price_requires_review(self):
        """A catalogue part cannot obtain a price from another part or a default."""
        result = match_items([self._item()], self.catalogue, {}, TODAY)[0]
        self.assertEqual(result.status, "needs review")
        self.assertIn("no approved price", result.reason)
        self.assertIsNone(result.price)

    def test_unclear_units_and_unit_aliases(self):
        """Ambiguous units block pricing; explicit each aliases remain acceptable."""
        for unit in ("", "box?", "m"):
            with self.subTest(unit=unit):
                result = match_items([self._item(unit=unit)], self.catalogue, self.prices, TODAY)[0]
                self.assertEqual(result.status, "needs review")
                self.assertIsNone(result.price)
        for unit in ("pcs", "each", "EA"):
            self.assertEqual(match_items([self._item(unit=unit)],
                                         self.catalogue, self.prices, TODAY)[0].status, "priced")

    def test_normalized_collision_requires_review(self):
        """Normalisation collisions do not pick a catalogue part automatically."""
        catalogue = dict(self.catalogue)
        catalogue["TRM101"] = Part("TRM101", "Different terminal block", "ea")
        result = match_items([self._item("trm 101")], catalogue, self.prices, TODAY)[0]
        self.assertEqual(result.status, "needs review")
        self.assertIsNone(result.price)
        self.assertEqual(len(result.candidates), 2)
        exact = match_items([self._item()], catalogue, self.prices, TODAY)[0]
        self.assertEqual(exact.status, "priced")

    def test_summary_counts(self):
        """Counts include every row and group expired prices into human review."""
        extraction, matches = self._matches()
        self.assertEqual(summarize(matches, extraction.unparsed),
                         {"priced": 7, "needs review": 3, "not in catalogue": 1,
                          "unparsed": 1, "price expired": 1, "parsed items": 11})

    def test_export_sheets_counts_provenance_and_unparsed(self):
        """The draft contains all sheets, accurate counts, hashes and source rows."""
        output = self._export()
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Quote draft", "Needs review", "Run info"])
        quote, review, info = book.worksheets
        self.assertEqual(quote.max_row, 13)
        self.assertEqual(review.max_row, 6)
        run_info = dict(info.iter_rows(min_row=2, values_only=True))
        self.assertEqual(run_info["priced"], 7)
        self.assertEqual(run_info["needs review"], 3)
        self.assertEqual(run_info["unparsed"], 1)
        self.assertEqual(run_info["Input file"], "rfq_a.pdf")
        self.assertEqual(run_info["Input SHA256"],
                         hashlib.sha256(self.files["rfq_a.pdf"].read_bytes()).hexdigest())
        self.assertEqual(run_info["Price evaluation date"], "2026-10-03")
        self.assertIn("Elapsed seconds before save", run_info)
        self.assertEqual(quote.cell(13, 10).value, "unparsed")
        self.assertIn("several", quote.cell(13, 14).value)
        self.assertIn("several", review.cell(review.max_row, 11).value)
        for row in quote.iter_rows(min_row=2, values_only=True):
            if row[9] != "priced":
                self.assertIsNone(row[6])
                self.assertIsNone(row[8])

    def test_export_same_second_never_overwrites(self):
        """Two identical timestamps create distinct intact run folders."""
        first = self._export()
        before = first.read_bytes()
        second = self._export()
        self.assertNotEqual(first.parent, second.parent)
        self.assertEqual(first.parent.name, "run-20261003-120000")
        self.assertEqual(second.parent.name, "run-20261003-120000-01")
        self.assertEqual(first.read_bytes(), before)

    def test_export_save_failure_leaves_no_partial_excel(self):
        """A write failure removes the new run while preserving an existing one."""
        first = self._export()
        before = first.read_bytes()
        def broken_save(workbook, destination):
            Path(destination).write_bytes(b"partial workbook")
            raise OSError("synthetic write failure")
        with patch("rfq_quote.export.Workbook.save", broken_save):
            with self.assertRaisesRegex(OSError, "synthetic write failure"):
                self._export()
        self.assertEqual(list((self.root / "out").iterdir()), [first.parent])
        self.assertEqual(first.read_bytes(), before)

    def test_export_rounding_invalid_zero_price_and_literal_text(self):
        """Money rounds half up; zero prices stay blank and strings stay literal."""
        items = [self._item(description="=HYPERLINK(fake)", qty="3"), self._item(qty="1")]
        prices = {"TRM-101": Price("TRM-101", Decimal("0.335"), "USD", date(2099, 1, 1))}
        matched = match_items(items[:1], self.catalogue, prices, TODAY)
        zero = {"TRM-101": Price("TRM-101", Decimal("0"), "USD", date(2099, 1, 1))}
        matched += match_items(items[1:], self.catalogue, zero, TODAY)
        output = export_quote(matched, [], self.files["rfq_a.pdf"], self.root / "out", TODAY, STARTED)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.active.cell(2, 9).value, 1.01)
        self.assertIsNone(book.active.cell(3, 7).value)
        self.assertIsNone(book.active.cell(3, 9).value)
        self.assertEqual(book.active.cell(3, 10).value, "invalid price")
        self.assertEqual(book.active.cell(2, 3).data_type, "s")

    def test_cli_main_errors_are_nonzero_without_excel(self):
        """Input errors return a nonzero code and do not leave an Excel draft."""
        args = [str(self.root / "missing.pdf"), "--catalogue", str(self.files["catalogue.csv"]),
                "--prices", str(self.files["price_list.csv"]), "--out-dir", str(self.root / "out")]
        stream = io.StringIO()
        with redirect_stderr(stream):
            self.assertEqual(main(args), 1)
        self.assertIn("Error:", stream.getvalue())
        self.assertEqual(list(self.root.rglob("*.xlsx")), [])

    def test_cli_main_success(self):
        """The public CLI function prints the four requested counts and output path."""
        args = [str(self.files["rfq_a.pdf"]), "--catalogue", str(self.files["catalogue.csv"]),
                "--prices", str(self.files["price_list.csv"]), "--today", TODAY.isoformat(),
                "--out-dir", str(self.root / "out")]
        stream = io.StringIO()
        with redirect_stdout(stream):
            self.assertEqual(main(args), 0)
        self.assertIn("priced: 7 / needs review: 3 / not in catalogue: 1 / unparsed: 1", stream.getvalue())
        self.assertEqual(len(list((self.root / "out").rglob("*.xlsx"))), 1)

    def test_cli_module_end_to_end_both_examples(self):
        """The documented module invocation runs both sample layouts to real Excel."""
        for name in ("rfq_a.pdf", "rfq_b.pdf"):
            with self.subTest(name=name):
                result = subprocess.run(
                    [sys.executable, "-m", "rfq_quote.cli", str(self.files[name]),
                     "--catalogue", str(self.files["catalogue.csv"]),
                     "--prices", str(self.files["price_list.csv"]),
                     "--today", TODAY.isoformat(), "--out-dir", str(self.root / name)],
                    cwd=ROOT, capture_output=True, text=True, check=False)
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertIn("priced: 7 / needs review: 3 / not in catalogue: 1 / unparsed: 1",
                              result.stdout)
                path = next((self.root / name).rglob("*.xlsx"))
                book = load_workbook(path)
                self.assertEqual(book.active.max_row, 13)
                book.close()

    def test_cli_module_failure_exit_code(self):
        """A subprocess observes a real nonzero exit status for a missing PDF."""
        result = subprocess.run(
            [sys.executable, "-m", "rfq_quote.cli", str(self.root / "missing.pdf"),
             "--catalogue", str(self.files["catalogue.csv"]),
             "--prices", str(self.files["price_list.csv"]), "--out-dir", str(self.root / "out")],
            cwd=ROOT, capture_output=True, text=True, check=False)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Error:", result.stderr)
        self.assertEqual(list(self.root.rglob("*.xlsx")), [])


if __name__ == "__main__":
    unittest.main()
