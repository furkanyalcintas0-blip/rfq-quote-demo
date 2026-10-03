"""Write a traceable Excel draft into a fresh timestamped run directory."""

import hashlib
import os
import shutil
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .extract import Unparsed
from .match import Match


def summarize(matches: list[Match], unparsed: list[Unparsed]) -> dict[str, int]:
    """Count outcomes; needs review includes expired/invalid prices and unit issues."""
    return {
        "priced": sum(m.status == "priced" for m in matches),
        "needs review": sum(m.status in ("needs review", "price expired", "invalid price")
                            for m in matches),
        "not in catalogue": sum(m.status == "not in catalogue" for m in matches),
        "unparsed": len(unparsed),
        "price expired": sum(m.status == "price expired" for m in matches),
        "parsed items": len(matches),
    }


def _append(sheet, values):
    sheet.append(values)
    # All user/source strings are literal text, even if they resemble formulas.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _style(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="254B62")
    for index, column in enumerate(sheet.columns, 1):
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 12), 65)


def export_quote(matches: list[Match], unparsed: list[Unparsed], input_path: str | Path,
                 out_dir: str | Path = "out", today: date | None = None,
                 started_at: datetime | None = None,
                 reference_paths: tuple[Path, ...] = ()) -> Path:
    """Save three sheets atomically, never overwriting an earlier run.

    Timestamps use UTC. Same-second collisions receive a numeric suffix.
    On save failure the new run directory is removed. Prices and line amounts
    appear only for priced rows; line amounts round half up to two decimals.
    Source text is written as literal text to prevent Excel formula injection.
    """
    started_at = started_at if started_at is not None else datetime.now(timezone.utc)
    if started_at.tzinfo is None:
        raise ValueError("started_at must include a timezone")
    today = today if today is not None else date.today()
    input_path = Path(input_path)
    counts = summarize(matches, unparsed)
    workbook = Workbook()
    quote = workbook.active
    quote.title = "Quote draft"
    _append(quote, ["Line", "Requested part", "Description", "Matched part", "Qty", "Unit",
                    "Unit price", "Currency", "Amount", "Status", "Note", "Page", "Source line",
                    "Source text"])
    review = workbook.create_sheet("Needs review")
    _append(review, ["Line", "Requested part", "Description", "Status", "Reason",
                     "Candidate 1", "Candidate 2", "Candidate 3", "Page", "Source line", "Source text"])
    for match in matches:
        item = match.item
        if match.status == "priced" and (match.price is None or match.part is None):
            raise ValueError("A priced row must have an approved price and a matched part")
        if match.status != "priced" and match.price is not None:
            raise ValueError("Review rows must not contain prices")
        price = match.price
        amount = ((item.qty * price.unit_price).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                  if price else None)
        _append(quote, [item.line_no, item.part_no, item.description,
                        match.part.part_no if match.part else None, float(item.qty), item.unit,
                        float(price.unit_price) if price else None, price.currency if price else None,
                        float(amount) if amount is not None else None, match.status, match.reason,
                        item.source.page, item.source.line, item.source.text])
        if match.status != "priced":
            candidates = [f"{c.part.part_no}: {c.part.description} [{c.part.unit}] "
                          f"(similarity {c.score:.2f})" for c in match.candidates]
            candidates += [None] * (3 - len(candidates))
            _append(review, [item.line_no, item.part_no, item.description, match.status, match.reason,
                             *candidates[:3], item.source.page, item.source.line, item.source.text])
    for row in unparsed:
        _append(quote, [None, None, None, None, None, None, None, None, None, "unparsed", row.reason,
                        row.source.page, row.source.line, row.source.text])
        _append(review, [None, None, None, "unparsed", row.reason, None, None, None,
                         row.source.page, row.source.line, row.source.text])
    info = workbook.create_sheet("Run info")
    _append(info, ["Field", "Value"])
    _append(info, ["Input file", input_path.name])
    _append(info, ["Input SHA256", hashlib.sha256(input_path.read_bytes()).hexdigest()])
    _append(info, ["Started at (UTC)", started_at.astimezone(timezone.utc).isoformat()])
    _append(info, ["Price evaluation date", today.isoformat()])
    for path in reference_paths:
        _append(info, [f"Reference file: {Path(path).name}",
                       hashlib.sha256(Path(path).read_bytes()).hexdigest()])
    for key, value in counts.items():
        _append(info, [key, value])
    _append(info, ["Review count definition",
                   "needs review includes price expired and invalid price; unparsed is separate"])
    _append(info, ["Output rows", len(matches) + len(unparsed)])
    _append(info, ["Draft notice", "Human review required; no substitutes selected; not a final quote"])
    for sheet in workbook:
        _style(sheet)
    for row in quote.iter_rows(min_row=2):
        row[4].number_format = "0.####"
        row[6].number_format = "0.00##"
        row[8].number_format = "0.00"
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = "run-" + started_at.astimezone(timezone.utc).strftime("%Y%m%d-%H%M%S")
    suffix = 0
    while True:
        run_dir = out_dir / (base if suffix == 0 else f"{base}-{suffix:02d}")
        try:
            run_dir.mkdir()
            break
        except FileExistsError:
            suffix += 1
    try:
        _append(info, ["Elapsed seconds before save",
                       round((datetime.now(timezone.utc) - started_at).total_seconds(), 3)])
        temporary = run_dir / ".quote-draft.tmp.xlsx"
        destination = run_dir / "quote-draft.xlsx"
        workbook.save(temporary)
        os.replace(temporary, destination)
        return destination
    except Exception:
        shutil.rmtree(run_dir)
        raise
    finally:
        workbook.close()
